"""Document parsing: PDF (PyMuPDF), DOCX, TXT + content-type detection."""
import io
import re

from app.utils.prompts import CONTENT_TYPE_KEYWORDS
from app.utils.validators import sanitize_text


def detect_content_type(text: str) -> str:
    lowered = text[:8000].lower()
    scores = {}
    for ctype, keywords in CONTENT_TYPE_KEYWORDS.items():
        scores[ctype] = sum(1 for kw in keywords if kw in lowered)
    best = max(scores, key=scores.get)
    return best if scores[best] > 0 else "lesson"


def extract_title(text: str, filename: str = "") -> str:
    ignore_prefixes = ("page", "|", "unit", "lesson", "chapter", "section", "course", "department")
    for line in text.splitlines()[:30]:
        line = line.strip().strip("#*- ").strip()
        line = re.sub(r"^[|\s\d\-_–—]+", "", line).strip()
        if not line:
            continue
        lower = line.lower()
        if any(lower.startswith(p) for p in ignore_prefixes):
            continue
        if re.search(r"\b\d{2}[A-Z]{2}\d{3}\b", line):
            continue
        if 6 <= len(line) <= 80 and len(line.split()) >= 2:
            return line.title()[:60]
    if filename:
        clean_fn = re.sub(r"^\d+[\s_-]*", "", filename)
        clean_fn = clean_fn.rsplit(".", 1)[0].replace("_", " ").replace("-", " ").strip()
        return clean_fn.title()[:60]
    return "Untitled analysis"


def parse_pdf(data: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as e:
        raise ValueError("PDF support is not installed on the server.") from e
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as e:
        raise ValueError("This PDF appears to be corrupted or password-protected.") from e
    parts = []
    try:
        if doc.page_count == 0:
            raise ValueError("This PDF has no readable pages.")
        for page in doc:
            parts.append(page.get_text("text"))
    finally:
        doc.close()
    text = sanitize_text("\n".join(parts))
    if len(text) < 50:
        raise ValueError("No readable text found in this PDF. It may be a scanned image — try a TXT or DOCX file.")
    return text


def parse_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise ValueError("DOCX support is not installed on the server.") from e
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise ValueError("This DOCX file appears to be corrupted.") from e
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    text = sanitize_text("\n".join(parts))
    if len(text) < 20:
        raise ValueError("This document is empty or has no readable text.")
    return text


def parse_txt(data: bytes) -> str:
    text = ""
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    text = sanitize_text(text)
    if len(text) < 20:
        raise ValueError("This file is empty or has no readable text.")
    return text


def parse_upload(filename: str, data: bytes) -> tuple[str, str]:
    """Returns (text, source_type). Raises ValueError with user-friendly message."""
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    if ext == "pdf":
        return parse_pdf(data), "pdf"
    if ext == "docx":
        return parse_docx(data), "docx"
    if ext in ("txt", "text", "md"):
        return parse_txt(data), "txt"
    raise ValueError(f"Unsupported file type '.{ext}'. Supported: PDF, DOCX, TXT.")
